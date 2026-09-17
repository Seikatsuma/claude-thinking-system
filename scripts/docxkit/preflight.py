#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Проверка вёрстки .docx перед отправкой человеку.

Ловит то, из-за чего документ приходит к заказчику кривым:
таблицы с автоподбором ширины, переполненные колонки, разорванные числа,
символы замены, лишние абзацы с разрывом страницы.

Запуск: preflight.py <файл.docx>   → печатает проблемы, код возврата 1 если они есть.
"""
import sys, re
from docx import Document
from docx.oxml.ns import qn

# ширина знака относительно кегля: Calibri ~0.48, DejaVu ~0.60. Берём пессимистичное.
CHAR_W = 0.60
NARROW_MM = 45              # колонка-поле: номер, сумма, имя
MAX_LINES_NARROW = 2        # в таких больше двух строк — уже плохо
MAX_LINES_WIDE = 6          # в описательных допустимо больше
TIGHT = 0.85                # длинное слово занимает больше этой доли ширины — впритык

TBLPR_ORDER = ['tblStyle','tblpPr','tblOverlap','bidiVisual','tblStyleRowBandSize','tblStyleColBandSize',
               'tblW','jc','tblCellSpacing','tblInd','tblBorders','shd','tblLayout','tblCellMar','tblLook',
               'tblCaption','tblDescription','tblPrChange']
TCPR_ORDER = ['cnfStyle','tcW','gridSpan','hMerge','vMerge','tcBorders','shd','noWrap','tcMar',
              'textDirection','tcFitText','vAlign','hideMark']
TRPR_ORDER = ['cnfStyle','divId','gridBefore','gridAfter','wBefore','wAfter','cantSplit','trHeight',
              'tblHeader','tblCellSpacing','jc','hidden','ins','del','trPrChange']

def _order_ok(el, order):
    idx = [order.index(c.tag.split('}')[1]) for c in el if c.tag.split('}')[1] in order]
    return idx == sorted(idx)

def check(path, verbose=True):
    d = Document(path)
    problems = []
    sec = d.sections[0]
    text_mm = (sec.page_width - sec.left_margin - sec.right_margin) / 36000.0

    for ti, t in enumerate(d.tables, 1):
        tblPr = t._tbl.tblPr
        lay = tblPr.find(qn('w:tblLayout'))
        if lay is None or lay.get(qn('w:type')) != 'fixed':
            problems.append(f'таблица {ti}: нет фиксированной раскладки — Word пересчитает ширины сам')
        if not _order_ok(tblPr, TBLPR_ORDER):
            problems.append(f'таблица {ti}: нарушен порядок элементов tblPr, файл может не открыться')

        grid = t._tbl.find(qn('w:tblGrid'))
        widths = [int(g.get(qn('w:w'))) / 56.7 for g in grid] if grid is not None else []
        if not widths:
            problems.append(f'таблица {ti}: не задана сетка колонок')
            continue
        total = sum(widths)
        if total > text_mm + 1:
            problems.append(f'таблица {ti}: шире полосы набора — {total:.0f} мм при {text_mm:.0f} мм')

        pad_mm = 2.3 * 2
        for ri, row in enumerate(t.rows, 1):
            if not _order_ok(row._tr.get_or_add_trPr(), TRPR_ORDER):
                problems.append(f'таблица {ti} строка {ri}: нарушен порядок trPr')
            cells, seen = [], set()
            for c in row.cells:
                if id(c._tc) not in seen:
                    seen.add(id(c._tc)); cells.append(c)
            if len(cells) != len(widths):
                continue
            for ci, (c, w) in enumerate(zip(cells, widths)):
                if c._tc.tcPr is not None and not _order_ok(c._tc.tcPr, TCPR_ORDER):
                    problems.append(f'таблица {ti} строка {ri} ячейка {ci+1}: нарушен порядок tcPr')
                usable = w - pad_mm
                for par in c.paragraphs:
                    txt = par.text
                    if not txt.strip():
                        continue
                    pt = max((r.font.size.pt for r in par.runs if r.font.size), default=11)
                    cw = pt * CHAR_W * 0.3528          # мм на знак
                    lines_n = len(txt) * cw / max(usable, 1)
                    limit = MAX_LINES_NARROW if w <= NARROW_MM else MAX_LINES_WIDE
                    if lines_n > limit:
                        problems.append(f'таблица {ti} строка {ri} колонка {ci+1} ({w:.0f} мм): '
                                        f'текст займёт ~{lines_n:.0f} строк — «{txt[:40]}»')
                    longest = max((len(x) for x in re.split(r'[\s\u00a0]+', txt)), default=0)
                    if longest * cw > usable:
                        problems.append(f'таблица {ti} строка {ri} колонка {ci+1}: '
                                        f'«{txt[:30]}» не влезает по ширине и будет разорвано')
                    elif longest * cw > usable * TIGHT:
                        problems.append(f'таблица {ti} строка {ri} колонка {ci+1}: «{txt[:30]}» '
                                        f'встаёт впритык — при другом шрифте перенесётся')

    blob = '\n'.join([p.text for p in d.paragraphs] +
                     [c.text for t in d.tables for r in t.rows for c in r.cells])
    if '�' in blob:
        problems.append('в тексте есть символы замены — сломана кодировка')
    if re.search(r'(?<=\d) (?=\d{3}\b)', blob):
        problems.append('числа разделены обычным пробелом — разорвутся при переносе строки '
                        '(нужен неразрывный)')
    for par in d.paragraphs:
        if par.paragraph_format.page_break_before and not par.text.strip():
            problems.append('пустой абзац с разрывом страницы — покажется пустым квадратом')

    if verbose:
        if problems:
            print(f'ПРОБЛЕМЫ ({len(problems)}):')
            for x in problems: print('  •', x)
        else:
            print('вёрстка в порядке')
    return problems

if __name__ == '__main__':
    sys.exit(1 if check(sys.argv[1]) else 0)
